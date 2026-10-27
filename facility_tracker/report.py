"""Turn the calculations into the two things a lender actually receives:
a compliance certificate (Markdown) and a supporting workbook (Excel)."""

from __future__ import annotations

from pathlib import Path

from .borrowing_base import BorrowingBase
from .covenants import CovenantResult
from .models import Quarter
from .stress import REVENUE_SHOCKS, VALUE_SHOCKS, StressCell


def _money(x: float) -> str:
    sign = "-" if x < 0 else ""
    return f"{sign}${abs(x) / 1e6:,.1f}M"


def _fmt(r: CovenantResult) -> tuple[str, str]:
    if r.unit == "x":
        return f"{r.actual:.2f}x", f"{'≥' if r.direction == 'min' else '≤'} {r.threshold:.2f}x"
    if r.unit == "%":
        return f"{r.actual:.1%}", f"{'≥' if r.direction == 'min' else '≤'} {r.threshold:.0%}"
    return _money(r.actual), f"{'≥' if r.direction == 'min' else '≤'} {_money(r.threshold)}"


def markdown(terms: dict, q: Quarter, bb: BorrowingBase, results: list[CovenantResult],
             stress: list[StressCell]) -> str:
    out = []
    out.append(f"# Compliance certificate: {q.period}")
    out.append("")
    out.append(f"**Borrower:** {terms['borrower']}  ")
    out.append(f"**Facility:** {terms['facility_name']}  ")
    out.append(f"**As of:** {bb.as_of.isoformat()}")
    out.append("")
    out.append("> Fictional borrower and illustrative terms, for demonstration only.")
    out.append("")

    out.append("## Summary")
    out.append("")
    status = "All covenants met" if all(r.passed for r in results) else "Covenant breach"
    bb_status = ("No borrowing base deficiency" if bb.deficiency == 0
                 else f"Borrowing base deficiency of {_money(bb.deficiency)}: mandatory prepayment required")
    out.append(f"- **Covenants:** {status}")
    out.append(f"- **Borrowing base:** {_money(bb.borrowing_base)} against {_money(bb.outstanding)} outstanding. {bb_status}")
    out.append(f"- **Availability:** {_money(max(bb.availability, 0))} undrawn")
    out.append("")

    out.append("## Financial covenants")
    out.append("")
    out.append("| Covenant | Actual | Requirement | Status | Headroom |")
    out.append("|---|---|---|---|---|")
    for r in results:
        actual, req = _fmt(r)
        out.append(f"| {r.name} | {actual} | {req} | {'Pass' if r.passed else '**Fail**'} | {r.headroom_note} |")
    out.append("")

    out.append("## Borrowing base")
    out.append("")
    out.append("| Cluster | GPU | Units | Customer | Collateral value | Eligible | Advance rate | Advance |")
    out.append("|---|---|---:|---|---:|---|---:|---:|")
    for line in bb.lines:
        c = line.cluster
        out.append(f"| {c.cluster_id} | {c.gpu_model} | {c.units:,} | {c.customer or '—'} | "
                   f"{_money(line.collateral_value)} | {line.reason} | {line.advance_rate:.0%} | {_money(line.advance_amount)} |")
    out.append(f"| **Total** | | | | **{_money(bb.total_collateral_value)}** | | | **{_money(bb.eligible_advance)}** |")
    out.append("")
    if bb.concentration_excess:
        for cust, amt in bb.concentration_excess.items():
            out.append(f"- Less concentration excess, {cust} (above {terms['concentration_limit_pct']:.0%} limit): {_money(-amt)}")
    else:
        out.append("- No customer above the concentration limit.")
    out.append(f"- **Borrowing base: {_money(bb.borrowing_base)}**")
    out.append("")

    out.append("## Stress test")
    out.append("")
    out.append("Availability under combined declines in GPU collateral value (rows) and quarterly revenue (columns). "
               "Cells marked ✗ breach at least one covenant. Negative availability means a borrowing base "
               "deficiency, which requires a mandatory prepayment even if every covenant is met.")
    out.append("")
    out.append("| Value decline | " + " | ".join(f"Revenue −{rs:.0%}" for rs in REVENUE_SHOCKS) + " |")
    out.append("|---|" + "---:|" * len(REVENUE_SHOCKS))
    for vs in VALUE_SHOCKS:
        row = [c for c in stress if c.value_shock == vs]
        cells = [f"{_money(c.availability)} {'✓' if c.all_covenants_pass else '✗'}" for c in row]
        out.append(f"| −{vs:.0%} | " + " | ".join(cells) + " |")
    out.append("")
    return "\n".join(out)


def excel(path: Path, terms: dict, q: Quarter, bb: BorrowingBase,
          results: list[CovenantResult], stress: list[StressCell]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DCE6E1")
    money = '#,##0;[Red](#,##0)'

    wb = Workbook()
    ws = wb.active
    ws.title = "Covenants"
    ws.append([f"{terms['borrower']} | {q.period} | as of {bb.as_of.isoformat()}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Fictional borrower, illustrative terms"])
    ws.append([])
    ws.append(["Covenant", "Actual", "Requirement", "Status", "Headroom"])
    for cell in ws[4]:
        cell.font, cell.fill = bold, header_fill
    for r in results:
        actual, req = _fmt(r)
        ws.append([r.name, actual, req, "Pass" if r.passed else "Fail", r.headroom_note])

    bs = wb.create_sheet("Borrowing base")
    bs.append(["Cluster", "GPU", "Units", "Customer", "Collateral value", "Eligibility", "Advance rate", "Advance"])
    for cell in bs[1]:
        cell.font, cell.fill = bold, header_fill
    for line in bb.lines:
        c = line.cluster
        bs.append([c.cluster_id, c.gpu_model, c.units, c.customer or "", round(line.collateral_value),
                   line.reason, line.advance_rate, round(line.advance_amount)])
    n = len(bb.lines) + 1
    bs.append(["Total", "", "", "", f"=SUM(E2:E{n})", "", "", f"=SUM(H2:H{n})"])
    for cust, amt in bb.concentration_excess.items():
        bs.append([f"Less concentration excess: {cust}", "", "", "", "", "", "", -round(amt)])
    last = bs.max_row
    bs.append(["Borrowing base", "", "", "", "", "", "", f"=SUM(H{n + 1}:H{last})"])
    bs.append(["Outstanding", "", "", "", "", "", "", bb.outstanding])
    bs.append(["Availability", "", "", "", "", "", "", f"=MIN({bb.commitment},H{last + 1})-H{last + 2}"])
    for row in bs.iter_rows(min_row=2, min_col=5, max_col=8):
        row[0].number_format = money
        row[2].number_format = "0%"
        row[3].number_format = money
    for r in range(n + 1, bs.max_row + 1):
        bs.cell(r, 1).font = bold

    st = wb.create_sheet("Stress")
    st.append(["Availability ($) by value decline (rows) and revenue decline (columns)"])
    st.append(["Value decline"] + [f"Revenue -{rs:.0%}" for rs in REVENUE_SHOCKS])
    for cell in st[2]:
        cell.font, cell.fill = bold, header_fill
    for vs in VALUE_SHOCKS:
        row = [c for c in stress if c.value_shock == vs]
        st.append([f"-{vs:.0%}"] + [round(c.availability) for c in row])
        for i, c in enumerate(row, start=2):
            cell = st.cell(st.max_row, i)
            cell.number_format = money
            if not c.all_covenants_pass:
                cell.fill = PatternFill("solid", fgColor="F4CCCC")
    st.append([])
    st.append(["Red cells breach at least one covenant."])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 48)
    wb.save(path)
